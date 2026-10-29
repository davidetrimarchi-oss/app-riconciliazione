import streamlit as st
import pandas as pd
import re
import io
from openpyxl.styles import Alignment

st.set_page_config(page_title="Riconciliazione Bancaria", layout="wide")
st.title("Riconciliazione Bancaria & Matching MN06 🏦")

# ==========================================
# 1. FUNZIONI DI SUPPORTO & PARSING
# ==========================================
def calcola_similarita_dealer(str1, str2):
    if not str1 or not str2 or pd.isna(str1) or pd.isna(str2):
        return 0.0
    s1, s2 = str(str1).upper().strip(), str(str2).upper().strip()
    if s1 == s2: return 1.0

    for char in [".", ",", "-", "/", "_", ":", ";"]:
        s1 = s1.replace(char, " ")
        s2 = s2.replace(char, " ")

    s1 = re.sub(r"\bAUTO\s+", "AUTO", s1)
    s2 = re.sub(r"\bAUTO\s+", "AUTO", s2)
    s1 = re.sub(r"\bFLLI\s+", "FLLI", s1)
    s2 = re.sub(r"\bFLLI\s+", "FLLI", s2)

    stop_words = {"SRL", "SPA", "SNC", "SAS", "CONCESSIONARIA", "AUTOMOBILI", "ITALIA", "GROUP"}
    parole1 = {p for p in s1.split() if len(p) >= 2 and p not in stop_words}
    parole2 = [p for p in s2.split() if len(p) >= 2 and p not in stop_words]

    if not parole2: return 0.0
    return sum(1 for p in parole2 if p in parole1) / len(parole2)

def estrai_mittente_pulito(text_raw):
    if not text_raw or pd.isna(text_raw): return ""
    t = str(text_raw).replace("\xa0", " ").strip()

    m_mitt = re.search(r'\bMITT\.?\s*:\s*(.*?)(?=\s*\bBENEF|\s*\bIBAN|\s*\bNOTE:|\s*\bDATA ORDINE|$)', t, re.IGNORECASE)
    if m_mitt:
        mitt = m_mitt.group(1).strip()
    else:
        m_disp = re.search(r'\b(?:DISPOSTO DA|ORDINANTE)\s*:\s*(.*?)(?=\s*\bBENEF|\s*\bIBAN|\s*\bNOTE:|\s*\bDATA ORDINE|$)', t, re.IGNORECASE)
        if m_disp:
            mitt = m_disp.group(1).strip()
            mitt = re.sub(r'\s+\d{8,}.*$', '', mitt).strip()
        else:
            mitt = ""
    return re.sub(r'[\.\s,;:\-]+$', '', mitt).strip()

def estrai_dati_bonifico(text_f):
    if not text_f or pd.isna(text_f): return "", []

    text_originale = str(text_f).replace("\xa0", " ").replace("\r", " ").replace("\n", " ").replace("\t", " ").strip()
    mittente = estrai_mittente_pulito(text_originale)

    pos_stop = text_originale.lower().find("bonifico a vostro favore")
    text = text_originale[:pos_stop] if pos_stop > 0 else text_originale

    text_clean = re.sub(r'\bDEL\s+\d{1,2}[\/\.\s\-]\d{1,2}(?:[\/\.\s\-](?:19|20)?\d{2})?\b', ' ', text, flags=re.IGNORECASE)
    text_clean = re.sub(r'\b\d{1,2}[\/\.\-]\d{1,2}[\/\.\-](?:19|20)?\d{2}\b', ' ', text_clean)
    text_clean = re.sub(r'\b\d{4}-\d{2}-\d{2}(?:\s+\d{2}:\d{2}:\d{2}(?:\.\d+)?)?\b', ' ', text_clean)
    text_clean = re.sub(r'\b(?:COD\.?\s*DISP\.?|CRO)\s*:\s*\w+', ' ', text_clean, flags=re.IGNORECASE)
    
    # 1. FIX INTELLIGENTI PER TELAI SCRITTI MALE
    text_clean = re.sub(r'(?i)OMODA\s*[eE]?\s*JAECOO', ' ', text_clean)
    text_clean = re.sub(r'(?i)(OMODA|JAECOO|DR[356]?|EVO\d?)', ' ', text_clean)
    
    def fix_spaced_vins(match):
        part1, part2 = match.group(1), match.group(2)
        if len(part1) + len(part2) == 17: return part1 + part2
        return match.group(0)
    
    text_clean = re.sub(r'\b([A-Z0-9]{3,14})\s+([A-Z0-9]{3,14})\b', fix_spaced_vins, text_clean, flags=re.IGNORECASE)
    text_clean = re.sub(r'\b(?:HIT|OIT|JIT|SIT|ZIT|AI|NIT)[A-Z0-9]*\b', ' ', text_clean, flags=re.IGNORECASE)
    text_clean = re.sub(r'\b(FATTURA|FATT|TELAIO|TEL|FT|DOC)([0-9]{4,})\b', r'\1 \2', text_clean, flags=re.IGNORECASE)

    # 2. ESTRAZIONE VIN E CHASSIS
    vin_list = []
    for v in re.findall(r'\b[A-HJ-NPR-Z0-9]{17}\b', text_clean, flags=re.IGNORECASE):
        v_upper = v.upper()
        if any(c.isalpha() for c in v_upper) and any(c.isdigit() for c in v_upper) and v_upper not in vin_list:
            vin_list.append(v_upper)

    text_for_chassis = text_clean
    boundaries = r'\b(PAG\.?|PAGAMENTO|SALDO|BONIFICO|CASH|ANTICIPO|ANTICIPATO|ACCONTO|RITIRO|TRADE\s*IN|FATT\.?|FATTURA|FAT\.?|FT\.?|DOC\.?|DOCUMENTO|NR\.?|N\.?|ORDINE|ORD\.?|CONTRATTO|COMMESSA|RIF\.?|RIFERIMENTO|DEL|DATA|IL|AL|MITT|BENEF|BIC|IBAN|CRO|EVO\d?|DR\d?|OMODA|JAECOO|COLLECTION|EURASIA|VETTURE|NUOVE|USATE|KM\s*0|KM0)(?=[0-9\W]|$)'
    text_for_chassis = re.sub(boundaries, '|', text_for_chassis, flags=re.IGNORECASE)

    chassis_list_temp = []
    parole_telai = r"TELAIO|TELAI|CHASSIS|VIN|TL|VEICOLO|TEL|TG|T\."
    
    for m in re.finditer(rf'\b(?:{parole_telai})\s*[:\.\-]?\s*([A-Z0-9]{{5,17}})', text_for_chassis, re.IGNORECASE):
        val = m.group(1).upper()
        if 8 < len(val) < 17: val = val[-8:]
        if len(val) < 17 and re.match(r'^[A-Z]{3,}', val): continue
        if len(val) == 17 and any(c.isalpha() for c in val) and any(c.isdigit() for c in val):
            if val not in vin_list: vin_list.append(val)
        elif any(c.isdigit() for c in val):
            if val not in chassis_list_temp: chassis_list_temp.append(val)

    for blocco in re.findall(rf'\b(?:{parole_telai})\s*[:\.\-]?\s*([A-Z0-9\s\+\,]+)', text_for_chassis, re.IGNORECASE):
        for p in re.split(r'[\+\,\s]+', blocco):
            token = p.strip().upper()
            if 8 < len(token) < 17: token = token[-8:]
            if 5 <= len(token) <= 17 and any(c.isdigit() for c in token):
                if len(token) < 17 and re.match(r'^[A-Z]{3,}', token): continue
                if len(token) == 17 and any(c.isalpha() for c in token):
                    if token not in vin_list: vin_list.append(token)
                elif token not in chassis_list_temp: chassis_list_temp.append(token)

    for match in re.finditer(r'\b([A-Z]{2,3}[0-9]{5,6}(?:-[A-Z]{2,3}[0-9]{5,6})*)\b', text_clean, re.IGNORECASE):
        potenziali_chassis = match.group(1).upper().split('-')
        for pc in potenziali_chassis:
            if len(pc) >= 6 and len(pc) <= 9 and any(c.isdigit() for c in pc):
                if pc not in chassis_list_temp and pc not in vin_list:
                    chassis_list_temp.append(pc)

    chassis_list = []
    chassis_con_lettera = [c for c in chassis_list_temp if c and c[0].isalpha()]
    for c in chassis_list_temp:
        if c.isdigit() and any(c in cl for cl in chassis_con_lettera): continue 
        if c not in chassis_list: chassis_list.append(c)

    # 3. ESTRAZIONE FATTURE 
    ordini_fatture = []
    parole_ordini = r"SALDO\s+FT|SALDO\s+FATTURA|SALDO\s+FATT|FATTURA\s+N|FATTURA|FATT|FATTURE|FT\s+NR|FT\s+N|FTIT\d*|FTRA|FT|INVOICE|INVOICES|INV|DOC\.?\s*N|DOC|DOCUMENTO|NR\s*DOC|CONTRATTO|CONTRATTI\s*N|CONTR|ORDINE|ORD|COMMESSA|RIF\.?|RIFERIMENTO"
    stop_tokens = {"DEL", "DATA", "IL", "AL", "PAG", "SAL", "SCAD", "BON", "CASH", "MITT", "BENEF", "BIC", "ORDINE", "DR3", "DR5", "EVO", "EVO5", "EVO6", "OMODA", "JAECOO", "VEICOLI"}

    # NUOVO: Aggiunto il blocco (?:N\.?\s*|NR\.?\s*|NUM\.?\s*)? per superare le abbreviazioni col punto (es: "N.")
    for match in re.finditer(rf'\b(?:{parole_ordini})\s*(?:N\.?\s*|NR\.?\s*|NUM\.?\s*)?[:\.\-°/]*\s*([A-Z0-9\/\-\+\,\s\(\)]+)', text_clean, re.IGNORECASE):
        for pezzo in re.split(r'[\+\/\,\s\(\)]+', match.group(1).strip().upper()):
            doc_val = re.sub(r'^(?:IT|NR|N)[\s\.\-]*', '', pezzo.strip())
            doc_val = re.sub(r'^(FATTURA|FATT|TELAIO|TEL|FT|DOC)\s*', '', doc_val)
            doc_val = re.sub(r'TELAIO$', '', doc_val)
            doc_val = re.sub(r'[\.\-\/]$', '', doc_val)
            
            for sub_doc in doc_val.split('-'):
                sub_doc = sub_doc.strip()
                if not sub_doc: continue
                if len(sub_doc) >= 16 and not sub_doc.isdigit(): continue
                is_invalid_short_code = bool(re.match(r'^[A-Z]{1,2}[0-9]{1,2}$', sub_doc))
                
                if len(sub_doc) >= 3 and any(c.isdigit() for c in sub_doc) and not is_invalid_short_code:
                    if sub_doc not in stop_tokens and sub_doc not in vin_list:
                        if not (sub_doc.isdigit() and len(sub_doc) == 2) and sub_doc not in ordini_fatture:
                            ordini_fatture.append(sub_doc)

    for match in re.finditer(r'\b(\d{6,8})\b', text_clean):
        doc_val = match.group(1)
        if doc_val not in stop_tokens and doc_val not in vin_list and doc_val not in chassis_list:
            if doc_val not in ordini_fatture:
                ordini_fatture.append(doc_val)

    # 4. ABBINAMENTO RECORD
    risultati = []
    veicoli = []
    
    if len(vin_list) > 0:
        for v in vin_list:
            chas = v[-7:]
            if not chas[0].isalpha() and len(v) >= 8:
                chas_8 = v[-8:]
                if chas_8[0].isalpha(): chas = chas_8
            veicoli.append({"VIN": v, "CHASSIS": chas})
    elif len(chassis_list) > 0:
        for c in chassis_list: veicoli.append({"VIN": "", "CHASSIS": c})
    else: veicoli.append({"VIN": "", "CHASSIS": ""})

    if len(veicoli) == 1 and len(ordini_fatture) > 1:
        for o in ordini_fatture: risultati.append({"VIN": veicoli[0]["VIN"], "CHASSIS": veicoli[0]["CHASSIS"], "Numero d'ordine / Fattura": o})
    elif len(ordini_fatture) <= 1 and len(veicoli) > 1:
        o_val = ordini_fatture[0] if ordini_fatture else ""
        for v in veicoli: risultati.append({"VIN": v["VIN"], "CHASSIS": v["CHASSIS"], "Numero d'ordine / Fattura": o_val})
    else:
        for i in range(max(len(veicoli), len(ordini_fatture), 1)):
            v = veicoli[i] if i < len(veicoli) else veicoli[-1]
            o = ordini_fatture[i] if i < len(ordini_fatture) else (ordini_fatture[-1] if ordini_fatture else "")
            risultati.append({"VIN": v["VIN"], "CHASSIS": v["CHASSIS"], "Numero d'ordine / Fattura": o})

    # Filtro anti-duplicati rigido per telaio
    risultati_unici = []
    visti_veicoli = set()
    
    for r in risultati:
        vin_val = r.get("VIN", "")
        chas_val = r.get("CHASSIS", "")
        
        if vin_val or chas_val:
            chiave_veicolo = (vin_val, chas_val)
            if chiave_veicolo in visti_veicoli:
                continue # Scarta la riga, il telaio è già presente
            visti_veicoli.add(chiave_veicolo)
            
        if r not in risultati_unici:
            risultati_unici.append(r)

    return mittente, risultati_unici

def recupera_vin_da_mn06(chassis_val, df_mn06, col_vin):
    if not chassis_val or pd.isna(chassis_val) or not col_vin: return ""

    if "_clean_chassis" in df_mn06.columns:
        cond_exact = df_mn06["_clean_chassis"] == chassis_val
        if cond_exact.any(): return str(df_mn06[cond_exact].iloc[0][col_vin]).strip()
        cond_ends = df_mn06["_clean_chassis"].str.endswith(chassis_val, na=False)
        if cond_ends.any(): return str(df_mn06[cond_ends].iloc[0][col_vin]).strip()

    if "_clean_vin" in df_mn06.columns:
        cond_vin_ends = df_mn06["_clean_vin"].str.endswith(chassis_val, na=False)
        if cond_vin_ends.any(): return str(df_mn06[cond_vin_ends].iloc[0]["_clean_vin"]).strip()

    return ""

def converti_importo(val):
    if pd.isna(val) or val is None: return 0.0
    if isinstance(val, (int, float)): return float(val)
    s = str(val).replace("€", "").replace("\xa0", "").strip()
    if "," in s and "." in s: s = s.replace(".", "").replace(",", ".")
    elif "," in s: s = s.replace(",", ".")
    try: return float(s)
    except ValueError: return 0.0

# ==========================================
# 2. INTERFACCIA STREAMLIT E MERGE EXCEL
# ==========================================
col1, col2 = st.columns(2)
with col1:
    file_banca = st.file_uploader("1. File Banca (TempoReale)", type=["xlsx", "xls", "csv"], key="banca")
with col2:
    file_mn06 = st.file_uploader("2. File Esposizione (MN06)", type=["xlsx", "xls", "csv"], key="mn06")

if file_banca and file_mn06:
    file_mn06.seek(0)
    preview_mn06 = pd.read_excel(file_mn06, nrows=5, header=None) if not file_mn06.name.endswith(".csv") else pd.read_csv(file_mn06, nrows=5, header=None)
    header_idx = 0
    for r_idx, r_vals in preview_mn06.iterrows():
        row_str = " ".join([str(v).lower() for v in r_vals if pd.notna(v)])
        if "chassis" in row_str or "vehicle identification" in row_str or "dealer" in row_str:
            header_idx = r_idx
            break

    file_mn06.seek(0)
    df_mn06 = pd.read_excel(file_mn06, header=header_idx) if not file_mn06.name.endswith(".csv") else pd.read_csv(file_mn06, header=header_idx)

    # MAPPATURA COLONNE CON PRIORITA'
    col_map = {}
    colonne_mn06 = list(df_mn06.columns)

    for col in colonne_mn06:
        c_clean = str(col).lower().replace("\n", " ").replace("\r", " ")
        c_clean = re.sub(r"\s+", " ", c_clean).strip()

        if "vin" not in col_map and ("vehicle identification number" in c_clean or c_clean == "vin"):
            col_map["vin"] = col
        if "chassis" not in col_map and "chassis" in c_clean:
            col_map["chassis"] = col
        if "dealer" not in col_map and ("dealer code text" in c_clean or ("dealer" in c_clean and "name" in c_clean) or c_clean == "dealer"):
            col_map["dealer"] = col
        if "gross" not in col_map and any(k in c_clean for k in ["gross", "outstanding", "outst", "esposizione", "residuo", "lordo", "balance", "open amount"]):
            col_map["gross"] = col

    for kw in ["invoice document number", "invoice", "fattura", "numero d'ordine", "document number", "ordine", "order", "doc", "rif"]:
        for col in colonne_mn06:
            if "inv" not in col_map:
                c_clean = str(col).lower().replace("\n", " ").replace("\r", " ").strip()
                if kw in c_clean:
                    col_map["inv"] = col
        if "inv" in col_map:
            break

    colonna_gross_selezionata = col_map.get("gross", None)
    if not colonna_gross_selezionata:
        st.warning("⚠️ Esposizione (Gross) non trovata automaticamente.")
        colonna_gross_selezionata = st.selectbox("Seleziona manualmente la colonna Gross Outstanding:", options=df_mn06.columns)
        col_map["gross"] = colonna_gross_selezionata

    if st.button("Avvia Elaborazione e Matching 🚀"):
        with st.spinner("Elaborazione in corso..."):
            file_banca.seek(0)
            df_banca_raw = pd.read_csv(file_banca) if file_banca.name.endswith(".csv") else pd.read_excel(file_banca, sheet_name=0, header=None)

            if "vin" in col_map: df_mn06["_clean_vin"] = df_mn06[col_map["vin"]].astype(str).str.strip().str.upper()
            else: df_mn06["_clean_vin"] = ""
            if "chassis" in col_map: df_mn06["_clean_chassis"] = df_mn06[col_map["chassis"]].astype(str).str.strip().str.upper()
            else: df_mn06["_clean_chassis"] = ""

            if "inv" in col_map: df_mn06["_clean_inv"] = df_mn06[col_map["inv"]].astype(str).str.replace(r'\.0$', '', regex=True).str.strip().str.upper()
            else: df_mn06["_clean_inv"] = ""

            rows_estratti = []
            start_row = 8 if not file_banca.name.endswith(".csv") else 0

            for idx in range(start_row, len(df_banca_raw)):
                row_data = df_banca_raw.iloc[idx]
                descrizione = row_data[5] if len(row_data) > 5 else ""
                importo = converti_importo(row_data[3]) if len(row_data) > 3 else 0.0

                if pd.isna(descrizione) or str(descrizione).strip() == "" or str(descrizione).strip() == "Descrizioni Aggiuntive":
                    continue

                mittente, record_veicoli = estrai_dati_bonifico(descrizione)

                for i, rec in enumerate(record_veicoli):
                    v_val = rec["VIN"]
                    c_val = rec["CHASSIS"]

                    if not v_val and c_val and "vin" in col_map:
                        v_trovato = recupera_vin_da_mn06(c_val, df_mn06, col_map["vin"])
                        if v_trovato: v_val = v_trovato

                    rows_estratti.append({
                        "Testo originale": str(descrizione),
                        "Mittente": mittente,
                        "VIN": v_val,
                        "CHASSIS": c_val,
                        "Numero d'ordine / Fattura": rec["Numero d'ordine / Fattura"],
                        "Importo Intesa": importo,
                        "Riga Orig.": idx + 1
                    })

            df_dest = pd.DataFrame(rows_estratti)

            # MATCHING CON MN06
            esiti, dealers, gross_list, righe_mn06 = [], [], [], []
            for idx_dest, r in df_dest.iterrows():
                vin_target = str(r["VIN"]).strip()
                chas_target = str(r["CHASSIS"]).strip()
                inv_target = str(r["Numero d'ordine / Fattura"]).strip()
                mitt_target = str(r["Mittente"]).strip()

                matches = []
                if vin_target or chas_target or inv_target:
                    cond = pd.Series(False, index=df_mn06.index)
                    if vin_target and "_clean_vin" in df_mn06: cond |= df_mn06["_clean_vin"].str.contains(vin_target, regex=False, na=False)
                    if chas_target and "_clean_chassis" in df_mn06: cond |= df_mn06["_clean_chassis"].str.contains(chas_target, regex=False, na=False)
                    if inv_target and "_clean_inv" in df_mn06: cond |= df_mn06["_clean_inv"].str.contains(inv_target, regex=False, na=False)
                    matches = df_mn06[cond].index.tolist()

                match_count = len(matches)
                if match_count > 0:
                    best_sim, best_idx, best_dealer = -1.0, matches[0], ""
                    dealer_col = col_map.get("dealer", None)

                    for m_idx in matches:
                        dealer_val = df_mn06.at[m_idx, dealer_col] if dealer_col else ""
                        sim = calcola_similarita_dealer(mitt_target, dealer_val)
                        if sim > best_sim:
                            best_sim, best_idx, best_dealer = sim, m_idx, str(dealer_val)

                    if not vin_target and "_clean_vin" in df_mn06:
                        v_mn06 = str(df_mn06.at[best_idx, "_clean_vin"]).strip()
                        if v_mn06 and v_mn06 != "NAN": df_dest.at[idx_dest, "VIN"] = v_mn06

                    if not chas_target and "_clean_chassis" in df_mn06:
                        c_mn06 = str(df_mn06.at[best_idx, "_clean_chassis"]).strip()
                        if c_mn06 and c_mn06 != "NAN": df_dest.at[idx_dest, "CHASSIS"] = c_mn06

                    if match_count == 1:
                        esito = "Trovato" if best_sim >= 0.6 else "Trovato - Solo Codice"
                    else:
                        esito = "Trovato (Risolto Fuzzy)" if best_sim >= 0.6 else "Verificare (Multipli su MN06)"

                    gross_val = converti_importo(df_mn06.at[best_idx, colonna_gross_selezionata])
                    esiti.append(esito); dealers.append(best_dealer); gross_list.append(gross_val); righe_mn06.append(best_idx + header_idx + 2)
                else:
                    esiti.append("Nessun Match"); dealers.append(""); gross_list.append(0.0); righe_mn06.append("")

            df_dest["Gross Outstanding MN06"] = gross_list
            df_dest["Dealer Name MN06"] = dealers
            df_dest["Esito Match MN06"] = esiti
            df_dest["Riga MN06"] = righe_mn06

            # VALIDAZIONE DELLE SOMME PER GRUPPO
            for riga_orig in df_dest["Riga Orig."].unique():
                mask = df_dest["Riga Orig."] == riga_orig
                group = df_dest[mask]

                importo_banca = group.iloc[0]["Importo Intesa"]
                somma_mn06 = group["Gross Outstanding MN06"].sum()
                is_sum_match = abs(importo_banca - somma_mn06) <= 0.05

                for i in group.index:
                    attuale_esito = df_dest.at[i, "Esito Match MN06"]
                    if attuale_esito == "Nessun Match": continue

                    if is_sum_match:
                        if "Trovato (Risolto Fuzzy)" in attuale_esito: df_dest.at[i, "Esito Match MN06"] = "OK - Match Totale (Risolto Fuzzy)"
                        elif "Solo Codice" in attuale_esito: df_dest.at[i, "Esito Match MN06"] = "OK - Match Totale (Dealer < 60%)"
                        elif "Multipli" in attuale_esito: df_dest.at[i, "Esito Match MN06"] = "OK - Match Totale (Scelto tra multipli)"
                        else: df_dest.at[i, "Esito Match MN06"] = "OK - Match Totale"
                    else:
                        diff = importo_banca - somma_mn06
                        if somma_mn06 == 0: df_dest.at[i, "Esito Match MN06"] = "Verificare - Nessun Importo Su MN06"
                        else: df_dest.at[i, "Esito Match MN06"] = f"Verificare - Somma Errata (Diff: {diff:,.2f}€)"

            st.success("Elaborazione completata con successo!")

            final_esiti = df_dest["Esito Match MN06"].tolist()
            m1, m2, m3, m4 = st.columns(4)
            m1.metric("Righe Generate", len(df_dest))
            m2.metric("Match Totali OK", sum(1 for e in final_esiti if "OK - Match" in e))
            m3.metric("Da Verificare", sum(1 for e in final_esiti if "Verificare" in e or "Solo Codice" in e))
            m4.metric("Nessun Match", sum(1 for e in final_esiti if e == "Nessun Match"))

            st.dataframe(df_dest.style.format({"Importo Intesa": "€ {:,.2f}", "Gross Outstanding MN06": "€ {:,.2f}"}), use_container_width=True)

            # ==========================================
            # CREAZIONE EXCEL CON CELLE UNITE FISICAMENTE (MERGE)
            # ==========================================
            buffer = io.BytesIO()
            with pd.ExcelWriter(buffer, engine='openpyxl') as writer:
                df_dest.to_excel(writer, index=False, sheet_name="Analisi_Estratti")
                workbook = writer.book
                worksheet = writer.sheets["Analisi_Estratti"]

                colonne_da_unire = [1, 2, 6, 7]
                riga_corrente_orig = None
                start_riga = 2

                for riga in range(2, len(df_dest) + 3):
                    valore_riga_orig = worksheet.cell(row=riga, column=7).value if riga <= len(df_dest) + 1 else None

                    if valore_riga_orig != riga_corrente_orig:
                        if riga_corrente_orig is not None:
                            end_riga = riga - 1
                            if end_riga > start_riga:
                                for col_idx in colonne_da_unire:
                                    worksheet.merge_cells(start_row=start_riga, start_column=col_idx, end_row=end_riga, end_column=col_idx)
                                    cella_unita = worksheet.cell(row=start_riga, column=col_idx)
                                    cella_unita.alignment = Alignment(vertical="center")

                        riga_corrente_orig = valore_riga_orig
                        start_riga = riga

                for cell in worksheet["F"]: cell.number_format = '€ #,##0.00'
                for cell in worksheet["H"]: cell.number_format = '€ #,##0.00'

            st.download_button(label="📥 Scarica Report Finale in Excel", data=buffer.getvalue(), file_name="Analisi_Estratti_Elaborato.xlsx", mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")